# -*- coding: utf-8 -*-
import sys, os
sys.stdout.reconfigure(encoding='utf-8')

from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn, nsdecls
from docx.oxml import OxmlElement, parse_xml
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image, ImageDraw, ImageFont
import copy

filepath = r'C:\Users\888\Desktop\基于结构混淆度引导的自适应随机分支图检索隐私保护方法.docx'
output_path = r'C:\Users\888\Desktop\基于结构混淆度引导的自适应随机分支图检索隐私保护方法_已修改.docx'

doc = Document(filepath)

# ============================================================
# Helper functions
# ============================================================

def set_run_font(run, font_name='宋体', size=Pt(16), bold=True):
    run.font.name = font_name
    run.font.size = size
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rFonts = rpr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rpr.insert(0, rFonts)
    rFonts.set(qn('w:eastAsia'), font_name)

def set_body_font(run):
    run.font.name = '宋体'
    run.font.size = Pt(12)
    rpr = run._element.get_or_add_rPr()
    rFonts = rpr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rpr.insert(0, rFonts)
    rFonts.set(qn('w:eastAsia'), '宋体')

def set_paragraph_spacing(paragraph, before=Pt(0), after=Pt(0), line_spacing=1.5):
    pf = paragraph.paragraph_format
    pf.space_before = before
    pf.space_after = after
    pf.line_spacing = line_spacing

def insert_element_after(target_elem, new_elem):
    parent = target_elem.getparent()
    parent.insert(list(parent).index(target_elem) + 1, new_elem)

# ============================================================
# MAIN HEADINGS
# ============================================================
MAIN_HEADINGS = [
    "方案名称", "技术领域", "背景技术", "发明内容", "发明目的",
    "技术方案", "有益效果", "附图说明", "具体实施方式", "实施例", "说明书摘要"
]
STEP_PREFIX = "步骤S"
SUB_METHODS = ["方式一：", "方式二：", "方式三：", "方式一", "方式二", "方式三"]

print("=== Phase 1: Formatting headings (宋体3号加粗) ===")

for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if not text:
        continue

    if text in MAIN_HEADINGS:
        for run in p.runs:
            set_run_font(run, size=Pt(16), bold=True)
        if text == "方案名称":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        print(f"  [{i:3d}] Section title: {text}")

    elif text.startswith(STEP_PREFIX) and len(text) > 3 and text[3].isdigit():
        for run in p.runs:
            set_run_font(run, size=Pt(16), bold=True)
        set_paragraph_spacing(p, before=Pt(6), after=Pt(3), line_spacing=1.5)
        print(f"  [{i:3d}] Step heading: {text}")

    elif text in SUB_METHODS:
        for run in p.runs:
            set_run_font(run, size=Pt(15), bold=True)
        print(f"  [{i:3d}] Method sub-heading: {text}")

print("\n=== Phase 2: Main title formatting ===")
# The main title is "一种基于结构混淆度引导的自适应随机分支图检索隐私保护方法" (P[1])
main_title_idx = 1
if len(doc.paragraphs) > main_title_idx:
    p = doc.paragraphs[main_title_idx]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        set_run_font(run, size=Pt(18), bold=True)
    set_paragraph_spacing(p, before=Pt(12), after=Pt(12), line_spacing=1.5)
    print(f"  Main title formatted: {p.text[:60]}...")

print("\n=== Phase 3: Page setup ===")
for section in doc.sections:
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.17)
    section.right_margin = Cm(3.17)
print("  A4 thesis margins set")

print("\n=== Phase 4: Body text formatting ===")
body_count = 0
for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if not text:
        continue
    if text in MAIN_HEADINGS:
        continue
    if text in SUB_METHODS:
        continue
    if text.startswith(STEP_PREFIX) and len(text) > 3 and text[3].isdigit():
        continue
    if i == 1:  # main title
        continue
    for run in p.runs:
        set_body_font(run)
    set_paragraph_spacing(p, before=Pt(0), after=Pt(0), line_spacing=1.5)
    body_count += 1
print(f"  Formatted {body_count} body paragraphs")

print("\n=== Phase 5: Inserting figure placeholder images ===")

figures = [
    ("图1", "本发明整体流程图"),
    ("图2", "结构混淆度计算示意图"),
    ("图3", "结构混淆度与保留概率映射关系图"),
    ("图4", "随机分支图检索执行过程示意图"),
    ("图5", "查询级随机化机制示意图"),
    ("图6", "本发明在Graph RAG系统中的部署结构图"),
    ("图7", "采用本发明前后图重构攻击效果对比示意图"),
]

# Find figure description paragraphs (the "图X为..." lines)
figure_targets = []
for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    for fig_num, fig_desc in figures:
        if text.startswith(fig_num):
            figure_targets.append((i, p, fig_num, fig_desc))
            break
figure_targets.sort(key=lambda x: x[0])
print(f"  Found {len(figure_targets)} figure description paragraphs")

# Create placeholder images
img_dir = os.path.join(os.path.dirname(output_path), "_figure_images")
os.makedirs(img_dir, exist_ok=True)

# Find a Chinese font
font_path = None
for fp in [
    r'C:\Windows\Fonts\simsun.ttc', r'C:\Windows\Fonts\simsun.ttf',
    r'C:\Windows\Fonts\msyh.ttc', r'C:\Windows\Fonts\msyh.ttf',
    r'C:\Windows\Fonts\msyhbd.ttc', r'C:\Windows\Fonts\yahei.ttf',
    r'C:\Windows\Fonts\deng.ttf', r'C:\Windows\Fonts\arial.ttf',
]:
    if os.path.exists(fp):
        font_path = fp
        break
print(f"  Using font: {font_path}")

# Process in reverse order so insertion doesn't shift paragraph indices
for rel_idx in range(len(figure_targets) - 1, -1, -1):
    para_idx, para, fig_num, fig_desc = figure_targets[rel_idx]

    # Create placeholder image
    img = Image.new('RGB', (680, 420), color=(255, 255, 255))
    draw = ImageDraw.Draw(img)
    draw.rectangle([3, 3, 677, 417], outline=(120, 120, 120), width=2)

    try:
        font_title = ImageFont.truetype(font_path, 28) if font_path else ImageFont.load_default()
        font_note = ImageFont.truetype(font_path, 18) if font_path else ImageFont.load_default()
    except:
        font_title = ImageFont.load_default()
        font_note = ImageFont.load_default()

    label = f"{fig_num}  {fig_desc}"
    bbox = draw.textbbox((0, 0), label, font=font_title)
    draw.text(((680 - (bbox[2]-bbox[0])) // 2, 140), label, fill=(60, 60, 60), font=font_title)
    note = "（请插入实际技术示意图）"
    bbox2 = draw.textbbox((0, 0), note, font=font_note)
    draw.text(((680 - (bbox2[2]-bbox2[0])) // 2, 200), note, fill=(180, 180, 180), font=font_note)

    # Simple graph icon
    draw.ellipse([270, 230, 290, 250], fill=(200, 200, 200))
    draw.line([280, 240, 350, 180], fill=(200, 200, 200), width=2)
    draw.ellipse([345, 170, 365, 190], fill=(200, 200, 200))
    draw.line([280, 240, 350, 300], fill=(200, 200, 200), width=2)
    draw.ellipse([345, 290, 365, 310], fill=(200, 200, 200))
    draw.line([280, 240, 200, 180], fill=(200, 200, 200), width=2)
    draw.ellipse([195, 170, 215, 190], fill=(200, 200, 200))
    draw.line([280, 240, 200, 300], fill=(200, 200, 200), width=2)
    draw.ellipse([195, 290, 215, 310], fill=(200, 200, 200))

    img_path = os.path.join(img_dir, f"figure_{rel_idx+1}.png")
    img.save(img_path)

    # ===== INSERT IMAGE - approach: add at end then move via lxml =====
    # Step 1: Add image paragraph at the end (python-docx handles relationship)
    img_para = doc.add_paragraph()
    img_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = img_para.add_run()
    run.add_picture(img_path, width=Inches(4.8))

    # Step 2: Add caption paragraph at the end
    cap_para = doc.add_paragraph()
    cap_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap_run = cap_para.add_run(f"{fig_num}　{fig_desc}")
    set_run_font(cap_run, font_name='宋体', size=Pt(10.5), bold=False)

    # Step 3: Move image paragraph after the figure description
    body = para._element.getparent()
    img_elem = img_para._element
    # Remove from end
    body.remove(img_elem)
    # Insert after description paragraph
    insert_element_after(para._element, img_elem)

    # Step 4: Move caption after image
    cap_elem = cap_para._element
    body.remove(cap_elem)
    insert_element_after(img_elem, cap_elem)

    print(f"  Inserted {fig_num} image after P[{para_idx}]")

print("\n=== Saving document ===")
doc.save(output_path)
print(f"Done! Saved to: {output_path}")

# ===== Verification =====
print("\n=== Verification ===")
import zipfile
with zipfile.ZipFile(output_path, 'r') as z:
    media = [f for f in z.namelist() if f.startswith('word/media/')]
    print(f"Images in document: {media}")
