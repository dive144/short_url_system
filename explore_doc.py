import sys
sys.stdout.reconfigure(encoding='utf-8')

import zipfile
from docx import Document

filepath = r'C:\Users\888\Desktop\基于结构混淆度引导的自适应随机分支图检索隐私保护方法.docx'
doc = Document(filepath)

print(f"Total paragraphs: {len(doc.paragraphs)}")
print()

# Print all paragraphs with style
for i, p in enumerate(doc.paragraphs):
    style = p.style.name if p.style else 'None'
    text = p.text.strip()
    if not text:
        continue
    # Check if this paragraph has any drawing/image
    has_drawing = bool(p._element.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing'))
    has_pict = bool(p._element.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pict'))
    img_mark = " [IMG]" if has_drawing or has_pict else ""
    runs_info = []
    for r in p.runs:
        font_name = r.font.name
        font_size = str(r.font.size) if r.font.size else 'inherit'
        bold = r.font.bold
        runs_info.append(f"font={font_name}, size={font_size}, bold={bold}")
    print(f"P[{i:3d}] {style:25s} {text[:100]}{img_mark}")
    if runs_info:
        for ri in runs_info[:1]:
            print(f"       runs: {ri}")

# Check images in zip
print("\n\n=== Images in media ===")
with zipfile.ZipFile(filepath, 'r') as z:
    media_files = sorted([f for f in z.namelist() if f.startswith('word/media/')])
    for mf in media_files:
        info = z.getinfo(mf)
        print(f"  {mf} ({info.file_size} bytes)")
